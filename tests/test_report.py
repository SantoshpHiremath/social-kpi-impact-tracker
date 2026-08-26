import os
import tempfile

from openpyxl import load_workbook

from src.data_quality import DataQualityFlag
from src.kpi_linkage import LinkedDepartmentSummary
from src.report import build_report


def make_linked(**overrides):
    defaults = dict(
        department="Engineering Strategy", quarter="2025-Q1",
        disability_representation_pct=6.0, accommodation_fulfillment_pct=75.0,
        training_completion_pct=80.0, on_time_delivery_pct=91.5,
        employee_engagement_score=4.1, voluntary_attrition_pct=5.0,
    )
    defaults.update(overrides)
    return LinkedDepartmentSummary(**defaults)


def test_report_file_created():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        result_path = build_report([make_linked()], [], output_path=path)
        assert result_path == path
        assert os.path.exists(path)


def test_report_has_expected_sheets():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        build_report([make_linked()], [], output_path=path)
        wb = load_workbook(path)
        assert set(wb.sheetnames) == {"Summary", "Accessibility & Performance", "Data Quality Flags"}


def test_linked_sheet_data_matches_input():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        linked = make_linked(department="Procurement", quarter="2025-Q3", on_time_delivery_pct=93.2)
        build_report([linked], [], output_path=path)
        wb = load_workbook(path)
        ws = wb["Accessibility & Performance"]
        row = [cell.value for cell in ws[2]]
        assert row[0] == "Procurement"
        assert row[1] == "2025-Q3"
        assert row[5] == 93.2


def test_none_fulfillment_written_as_na():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        linked = make_linked(accommodation_fulfillment_pct=None)
        build_report([linked], [], output_path=path)
        wb = load_workbook(path)
        ws = wb["Accessibility & Performance"]
        row = [cell.value for cell in ws[2]]
        assert row[3] == "n/a"


def test_none_training_written_as_not_reported():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        linked = make_linked(training_completion_pct=None)
        build_report([linked], [], output_path=path)
        wb = load_workbook(path)
        ws = wb["Accessibility & Performance"]
        row = [cell.value for cell in ws[2]]
        assert row[4] == "not reported"


def test_flags_sheet_has_flag_rows():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        flag = DataQualityFlag(department="Production", quarter="2026-Q1",
                                issue_type="accommodation_overcount", detail="test detail")
        build_report([make_linked()], [flag], output_path=path)
        wb = load_workbook(path)
        ws = wb["Data Quality Flags"]
        row = [cell.value for cell in ws[2]]
        assert row[0] == "Production"
        assert row[2] == "accommodation_overcount"


def test_summary_sheet_counts_correct():
    with tempfile.TemporaryDirectory() as tmpdir:
        path = os.path.join(tmpdir, "report.xlsx")
        linked = [make_linked(), make_linked(department="Procurement")]
        flags = [DataQualityFlag(department="X", quarter="2025-Q1", issue_type="t", detail="d")]
        build_report(linked, flags, output_path=path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        values = [row[1] for row in ws.iter_rows(min_row=3, max_row=4, values_only=True)]
        assert values[0] == 2  # linked records count
        assert values[1] == 1  # flags count
