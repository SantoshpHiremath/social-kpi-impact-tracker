"""
Writes a real, formatted .xlsx tracking report -- the deliverable a
CSR/social-KPI tracking tool actually needs to produce for leadership.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.data_quality import DataQualityFlag
from src.kpi_linkage import LinkedDepartmentSummary

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_header(ws, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22


def _write_summary_sheet(wb: Workbook, linked: list[LinkedDepartmentSummary], flags: list[DataQualityFlag]) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Social KPI & Departmental Impact Tracker"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Linked department/quarter records", len(linked)])
    ws.append(["Data-quality flags raised", len(flags)])
    quarters = sorted({r.quarter for r in linked})
    ws.append(["Quarters covered", f"{quarters[0]} to {quarters[-1]}" if quarters else "n/a"])
    ws.append(["Departments covered", len(sorted({r.department for r in linked}))])


def _write_linked_sheet(wb: Workbook, linked: list[LinkedDepartmentSummary]) -> None:
    ws = wb.create_sheet("Accessibility & Performance")
    headers = [
        "Department", "Quarter", "Disability Representation %",
        "Accommodation Fulfillment %", "Training Completion %",
        "On-Time Delivery %", "Employee Engagement (1-5)", "Voluntary Attrition %",
    ]
    _write_header(ws, headers)
    for row in linked:
        ws.append([
            row.department, row.quarter, row.disability_representation_pct,
            row.accommodation_fulfillment_pct if row.accommodation_fulfillment_pct is not None else "n/a",
            row.training_completion_pct if row.training_completion_pct is not None else "not reported",
            row.on_time_delivery_pct, row.employee_engagement_score, row.voluntary_attrition_pct,
        ])


def _write_flags_sheet(wb: Workbook, flags: list[DataQualityFlag]) -> None:
    ws = wb.create_sheet("Data Quality Flags")
    _write_header(ws, ["Department", "Quarter", "Issue Type", "Detail"])
    for f in flags:
        ws.append([f.department, f.quarter, f.issue_type, f.detail])


def build_report(
    linked: list[LinkedDepartmentSummary],
    flags: list[DataQualityFlag],
    output_path: str = "output/social_kpi_impact_report.xlsx",
) -> str:
    wb = Workbook()
    _write_summary_sheet(wb, linked, flags)
    _write_linked_sheet(wb, linked)
    _write_flags_sheet(wb, flags)
    wb.save(output_path)
    return output_path
